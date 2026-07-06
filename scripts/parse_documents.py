#!/usr/bin/env python3
"""Parse crawled public documents into JSONL documents and chunks.

This is intentionally pragmatic for the project MVP:
- parse common document types in src/data
- preserve institution/file metadata
- write a parse report for failures and skipped files
- create simple paragraph-aware chunks for the first retrieval index

Examples:
  python scripts/parse_documents.py --input src/data --output processed --workers 4
  python scripts/parse_documents.py --institution 금융감독원 --limit 20
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import logging
import os
import re
import struct
import sys
import traceback
import warnings
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".hwp",
    ".hwpx",
    ".docx",
    ".xlsx",
    ".xls",
    ".pptx",
}
SKIPPED_EXTENSIONS = {
    ".csv",
    ".doc",
    ".jpg",
    ".jsonl",
    ".mp4",
    ".png",
    ".zip",
}
DEFAULT_CHUNK_CHARS = 1800
DEFAULT_CHUNK_OVERLAP = 250

logging.getLogger("pypdf").setLevel(logging.ERROR)
logging.getLogger("pypdf._cmap").setLevel(logging.ERROR)
warnings.filterwarnings("ignore", module="pypdf")


@dataclass(frozen=True)
class ParseTask:
    path: Path
    data_root: Path
    output_root: Path
    chunk_chars: int
    chunk_overlap: int
    min_chars: int
    pdf_fallback: bool
    max_file_mb: float | None


def json_default(value: Any) -> str:
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def stable_id(value: str) -> str:
    return hashlib.sha1(value.encode("utf-8", errors="ignore")).hexdigest()[:16]


def to_posix(path: Path) -> str:
    return path.as_posix()


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f]", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def file_metadata(path: Path, data_root: Path) -> dict[str, Any]:
    rel = path.relative_to(data_root)
    parts = rel.parts
    institution = parts[0] if len(parts) > 1 else ""
    stat = path.stat()
    rel_posix = to_posix(rel)
    return {
        "doc_id": stable_id(rel_posix),
        "institution": institution,
        "source_path": to_posix(Path("src/data") / rel),
        "relative_path": rel_posix,
        "file_name": path.name,
        "extension": path.suffix.lower(),
        "size_bytes": stat.st_size,
        "modified_at": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
    }


def parse_pdf(path: Path, use_fallback: bool) -> tuple[str, dict[str, Any]]:
    try:
        import fitz

        page_texts: list[str] = []
        with fitz.open(str(path)) as document:
            page_count = document.page_count
            for index, page in enumerate(document, start=1):
                text = page.get_text("text") or ""
                if text.strip():
                    page_texts.append(f"[page {index}]\n{text}")
        text = clean_text("\n\n".join(page_texts))
        if text:
            return text, {"parser": "pymupdf", "page_count": page_count}
    except Exception:
        pass

    import pypdf

    reader = pypdf.PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            pass

    page_texts: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            page_texts.append(f"[page {index}]\n{text}")

    parser = "pypdf"
    text = "\n\n".join(page_texts)
    if use_fallback and len(clean_text(text)) < 50:
        try:
            import pdfplumber

            page_texts = []
            with pdfplumber.open(str(path)) as pdf:
                for index, page in enumerate(pdf.pages, start=1):
                    extracted = page.extract_text() or ""
                    if extracted.strip():
                        page_texts.append(f"[page {index}]\n{extracted}")
            fallback_text = "\n\n".join(page_texts)
            if len(fallback_text) > len(text):
                text = fallback_text
                parser = "pdfplumber"
        except Exception:
            pass

    return clean_text(text), {"parser": parser, "page_count": len(reader.pages)}


def hwp_is_compressed(ole: Any) -> bool:
    header = ole.openstream("FileHeader").read()
    if len(header) < 40:
        return False
    flags = struct.unpack_from("<I", header, 36)[0]
    return bool(flags & 0x01)


def sort_hwp_section_name(name: str) -> tuple[int, str]:
    match = re.search(r"Section(\d+)$", name)
    if match:
        return int(match.group(1)), name
    return 10**9, name


def decode_hwp_record_text(payload: bytes) -> str:
    text = payload.decode("utf-16le", errors="ignore")
    text = text.replace("\u0000", " ")
    text = text.replace("\u000d", "\n")
    return text


def extract_hwp_section_text(data: bytes) -> str:
    chunks: list[str] = []
    offset = 0
    while offset + 4 <= len(data):
        header = int.from_bytes(data[offset : offset + 4], "little")
        offset += 4
        tag_id = header & 0x3FF
        size = (header >> 20) & 0xFFF
        if size == 0xFFF:
            if offset + 4 > len(data):
                break
            size = int.from_bytes(data[offset : offset + 4], "little")
            offset += 4
        if size < 0 or offset + size > len(data):
            break
        payload = data[offset : offset + size]
        offset += size
        if tag_id == 67:  # HWPTAG_PARA_TEXT
            value = decode_hwp_record_text(payload)
            if value.strip():
                chunks.append(value)
    return "\n".join(chunks)


def parse_hwp(path: Path) -> tuple[str, dict[str, Any]]:
    try:
        import olefile
    except ImportError as exc:
        raise RuntimeError("Install olefile to parse .hwp files: python -m pip install olefile") from exc

    if not olefile.isOleFile(str(path)):
        if zipfile.is_zipfile(path):
            text, metadata = parse_hwpx(path)
            metadata["parser"] = "hwpx-xml-renamed-hwp"
            return text, metadata
        raise RuntimeError("Not an OLE HWP file")

    texts: list[str] = []
    section_count = 0
    with olefile.OleFileIO(str(path)) as ole:
        compressed = hwp_is_compressed(ole)
        streams = ["/".join(item) for item in ole.listdir(streams=True)]
        sections = sorted(
            [name for name in streams if name.startswith("BodyText/Section")],
            key=sort_hwp_section_name,
        )
        for section in sections:
            raw = ole.openstream(section).read()
            if compressed:
                import zlib

                try:
                    raw = zlib.decompress(raw, -15)
                except zlib.error:
                    raw = zlib.decompress(raw)
            section_text = extract_hwp_section_text(raw)
            if section_text.strip():
                section_count += 1
                texts.append(section_text)

        if not texts and "PrvText" in streams:
            raw = ole.openstream("PrvText").read()
            for encoding in ("utf-16le", "cp949", "utf-8"):
                try:
                    preview = raw.decode(encoding, errors="ignore")
                    if preview.strip():
                        texts.append(preview)
                        break
                except Exception:
                    continue

    return clean_text("\n\n".join(texts)), {
        "parser": "hwp5-ole",
        "section_count": section_count,
    }


def iter_xml_text_from_zip(path: Path, prefixes: tuple[str, ...]) -> Iterable[str]:
    with zipfile.ZipFile(path) as archive:
        names = sorted(
            name
            for name in archive.namelist()
            if name.lower().endswith(".xml") and name.startswith(prefixes)
        )
        for name in names:
            try:
                xml_bytes = archive.read(name)
                root = ElementTree.fromstring(xml_bytes)
                text = " ".join(part.strip() for part in root.itertext() if part and part.strip())
                if text:
                    yield html.unescape(text)
            except Exception:
                continue


def parse_hwpx(path: Path) -> tuple[str, dict[str, Any]]:
    texts = list(iter_xml_text_from_zip(path, ("Contents/", "BodyText/")))
    return clean_text("\n\n".join(texts)), {"parser": "hwpx-xml", "xml_parts": len(texts)}


def parse_docx(path: Path) -> tuple[str, dict[str, Any]]:
    try:
        import docx

        document = docx.Document(str(path))
        texts: list[str] = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                texts.append(paragraph.text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    texts.append("\t".join(cells))
        return clean_text("\n".join(texts)), {"parser": "python-docx"}
    except Exception:
        texts = list(iter_xml_text_from_zip(path, ("word/",)))
        return clean_text("\n\n".join(texts)), {"parser": "docx-xml", "xml_parts": len(texts)}


def parse_xlsx(path: Path) -> tuple[str, dict[str, Any]]:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    texts: list[str] = []
    row_count = 0
    for sheet in workbook.worksheets:
        texts.append(f"[sheet] {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value).strip() for value in row]
            while values and values[-1] == "":
                values.pop()
            if values:
                row_count += 1
                texts.append("\t".join(values))
    workbook.close()
    return clean_text("\n".join(texts)), {"parser": "openpyxl", "sheet_count": len(workbook.sheetnames), "row_count": row_count}


def parse_xls(path: Path) -> tuple[str, dict[str, Any]]:
    try:
        import pandas as pd

        sheets = pd.read_excel(str(path), sheet_name=None, dtype=str, header=None)
    except Exception as exc:
        raise RuntimeError(f"Unable to parse .xls; install xlrd or convert to .xlsx: {exc}") from exc

    texts: list[str] = []
    row_count = 0
    for name, frame in sheets.items():
        texts.append(f"[sheet] {name}")
        frame = frame.fillna("")
        for row in frame.astype(str).itertuples(index=False, name=None):
            values = [value.strip() for value in row]
            while values and values[-1] == "":
                values.pop()
            if values:
                row_count += 1
                texts.append("\t".join(values))
    return clean_text("\n".join(texts)), {"parser": "pandas-excel", "sheet_count": len(sheets), "row_count": row_count}


def parse_pptx(path: Path) -> tuple[str, dict[str, Any]]:
    texts = list(iter_xml_text_from_zip(path, ("ppt/slides/",)))
    return clean_text("\n\n".join(texts)), {"parser": "pptx-xml", "slide_count": len(texts)}


def parse_file(path: Path, pdf_fallback: bool) -> tuple[str, dict[str, Any]]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return parse_pdf(path, pdf_fallback)
    if ext == ".hwp":
        return parse_hwp(path)
    if ext == ".hwpx":
        return parse_hwpx(path)
    if ext == ".docx":
        return parse_docx(path)
    if ext == ".xlsx":
        return parse_xlsx(path)
    if ext == ".xls":
        return parse_xls(path)
    if ext == ".pptx":
        return parse_pptx(path)
    raise RuntimeError(f"Unsupported extension: {ext}")


def paragraph_chunks(text: str, max_chars: int, overlap: int) -> list[dict[str, Any]]:
    paragraphs = [part.strip() for part in re.split(r"\n{2,}", text) if part.strip()]
    chunks: list[dict[str, Any]] = []
    current: list[str] = []
    current_len = 0

    def flush() -> None:
        nonlocal current, current_len
        if not current:
            return
        chunk_text = "\n\n".join(current).strip()
        if chunk_text:
            chunks.append({"text": chunk_text, "char_count": len(chunk_text)})
        if overlap > 0 and chunk_text:
            tail = chunk_text[-overlap:]
            current = [tail]
            current_len = len(tail)
        else:
            current = []
            current_len = 0

    for paragraph in paragraphs:
        if len(paragraph) > max_chars:
            flush()
            start = 0
            step = max(1, max_chars - overlap)
            while start < len(paragraph):
                piece = paragraph[start : start + max_chars].strip()
                if piece:
                    chunks.append({"text": piece, "char_count": len(piece)})
                if start + max_chars >= len(paragraph):
                    break
                start += step
            continue

        next_len = current_len + len(paragraph) + (2 if current else 0)
        if current and next_len > max_chars:
            flush()
        current.append(paragraph)
        current_len += len(paragraph) + (2 if current_len else 0)

    flush()
    return chunks


def build_record(task: ParseTask) -> dict[str, Any]:
    metadata = file_metadata(task.path, task.data_root)
    ext = metadata["extension"]
    if task.max_file_mb and metadata["size_bytes"] > task.max_file_mb * 1024 * 1024:
        return {
            "metadata": metadata,
            "status": "deferred",
            "reason": f"over_max_file_mb:{task.max_file_mb}",
            "text": "",
            "chunks": [],
        }
    if ext in SKIPPED_EXTENSIONS:
        return {
            "metadata": metadata,
            "status": "skipped",
            "reason": f"skipped_extension:{ext}",
            "text": "",
            "chunks": [],
        }
    if ext not in SUPPORTED_EXTENSIONS:
        return {
            "metadata": metadata,
            "status": "skipped",
            "reason": f"unsupported_extension:{ext}",
            "text": "",
            "chunks": [],
        }

    try:
        text, extra = parse_file(task.path, task.pdf_fallback)
        metadata.update(extra)
        char_count = len(text)
        metadata["char_count"] = char_count
        if char_count < task.min_chars:
            return {
                "metadata": metadata,
                "status": "empty",
                "reason": f"below_min_chars:{char_count}",
                "text": text,
                "chunks": [],
            }
        chunks = paragraph_chunks(text, task.chunk_chars, task.chunk_overlap)
        return {
            "metadata": metadata,
            "status": "parsed",
            "reason": "",
            "text": text,
            "chunks": chunks,
        }
    except Exception as exc:
        metadata["char_count"] = 0
        return {
            "metadata": metadata,
            "status": "error",
            "reason": f"{type(exc).__name__}: {exc}",
            "traceback": traceback.format_exc(limit=5),
            "text": "",
            "chunks": [],
        }


def iter_input_files(data_root: Path, institutions: set[str] | None, extensions: set[str] | None) -> list[Path]:
    files: list[Path] = []
    for path in data_root.rglob("*"):
        if not path.is_file():
            continue
        rel = path.relative_to(data_root)
        institution = rel.parts[0] if len(rel.parts) > 1 else ""
        if institutions and institution not in institutions:
            continue
        if extensions and path.suffix.lower() not in extensions:
            continue
        files.append(path)
    return sorted(files, key=lambda item: to_posix(item.relative_to(data_root)))


def write_outputs(records: Iterable[dict[str, Any]], output_root: Path) -> dict[str, int]:
    output_root.mkdir(parents=True, exist_ok=True)
    docs_path = output_root / "documents.jsonl"
    chunks_path = output_root / "chunks.jsonl"
    report_path = output_root / "parse_report.csv"

    counts = {"parsed": 0, "empty": 0, "error": 0, "skipped": 0, "deferred": 0, "chunks": 0}
    with docs_path.open("w", encoding="utf-8", newline="\n") as docs_file, chunks_path.open(
        "w", encoding="utf-8", newline="\n"
    ) as chunks_file, report_path.open("w", encoding="utf-8-sig", newline="") as report_file:
        report = csv.DictWriter(
            report_file,
            fieldnames=[
                "status",
                "reason",
                "institution",
                "relative_path",
                "extension",
                "parser",
                "char_count",
                "chunk_count",
                "size_bytes",
            ],
        )
        report.writeheader()
        for record in records:
            metadata = record["metadata"]
            status = record["status"]
            chunks = record.get("chunks") or []
            counts[status] = counts.get(status, 0) + 1
            counts["chunks"] += len(chunks)

            if status in {"parsed", "empty"}:
                doc_record = {
                    **metadata,
                    "status": status,
                    "text": record["text"],
                }
                docs_file.write(json.dumps(doc_record, ensure_ascii=False, default=json_default) + "\n")

            if status == "parsed":
                for index, chunk in enumerate(chunks):
                    chunk_record = {
                        "chunk_id": f"{metadata['doc_id']}#{index:04d}",
                        "doc_id": metadata["doc_id"],
                        "chunk_index": index,
                        "text": chunk["text"],
                        "char_count": chunk["char_count"],
                        "metadata": {
                            key: metadata.get(key)
                            for key in (
                                "institution",
                                "source_path",
                                "relative_path",
                                "file_name",
                                "extension",
                                "parser",
                            )
                        },
                    }
                    chunks_file.write(json.dumps(chunk_record, ensure_ascii=False, default=json_default) + "\n")

            report.writerow(
                {
                    "status": status,
                    "reason": record.get("reason", ""),
                    "institution": metadata.get("institution", ""),
                    "relative_path": metadata.get("relative_path", ""),
                    "extension": metadata.get("extension", ""),
                    "parser": metadata.get("parser", ""),
                    "char_count": metadata.get("char_count", 0),
                    "chunk_count": len(chunks),
                    "size_bytes": metadata.get("size_bytes", 0),
                }
            )

    return counts


def iter_records(tasks: list[ParseTask], workers: int) -> Iterable[dict[str, Any]]:
    completed = 0
    total = len(tasks)
    if workers <= 1:
        for task in tasks:
            yield build_record(task)
            completed += 1
            if completed % 25 == 0 or completed == total:
                print(f"Parsed {completed}/{total}")
        return

    with ThreadPoolExecutor(max_workers=workers) as executor:
        futures = [executor.submit(build_record, task) for task in tasks]
        for future in as_completed(futures):
            yield future.result()
            completed += 1
            if completed % 25 == 0 or completed == total:
                print(f"Parsed {completed}/{total}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", default="src/data", help="Directory containing institution document folders.")
    parser.add_argument("--output", default="processed", help="Directory for documents.jsonl/chunks.jsonl/report.")
    parser.add_argument("--institution", action="append", help="Institution folder to parse. Can be repeated.")
    parser.add_argument("--extensions", nargs="*", help="Optional extension allowlist, e.g. .pdf .hwp")
    parser.add_argument("--limit", type=int, help="Limit number of files for smoke tests.")
    parser.add_argument("--workers", type=int, default=max(1, min(4, (os.cpu_count() or 2) - 1)))
    parser.add_argument("--chunk-chars", type=int, default=DEFAULT_CHUNK_CHARS)
    parser.add_argument("--chunk-overlap", type=int, default=DEFAULT_CHUNK_OVERLAP)
    parser.add_argument("--min-chars", type=int, default=20)
    parser.add_argument("--pdf-fallback", action="store_true", help="Use pdfplumber if pypdf extracts almost no text.")
    parser.add_argument("--max-file-mb", type=float, help="Defer files larger than this size in MB.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    data_root = Path(args.input).resolve()
    output_root = Path(args.output).resolve()
    institutions = set(args.institution) if args.institution else None
    extensions = {item.lower() if item.startswith(".") else f".{item.lower()}" for item in args.extensions or []}
    files = iter_input_files(data_root, institutions, extensions or None)
    if args.limit:
        files = files[: args.limit]

    if not data_root.exists():
        print(f"Input directory does not exist: {data_root}", file=sys.stderr)
        return 2
    if not files:
        print("No files matched.")
        return 0

    print(f"Parsing {len(files)} files from {data_root}")
    print(f"Output: {output_root}")
    print(f"Workers: {args.workers}")

    tasks = [
        ParseTask(
            path=path,
            data_root=data_root,
            output_root=output_root,
            chunk_chars=args.chunk_chars,
            chunk_overlap=args.chunk_overlap,
            min_chars=args.min_chars,
            pdf_fallback=args.pdf_fallback,
            max_file_mb=args.max_file_mb,
        )
        for path in files
    ]

    counts = write_outputs(iter_records(tasks, args.workers), output_root)
    summary_path = output_root / "parse_summary.json"
    summary = {
        "input": str(data_root),
        "output": str(output_root),
        "file_count": len(files),
        "counts": counts,
        "chunk_chars": args.chunk_chars,
        "chunk_overlap": args.chunk_overlap,
        "max_file_mb": args.max_file_mb,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
