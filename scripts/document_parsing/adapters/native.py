"""Native, low-cost adapters for HWP/HWPX, HTML, and Office documents."""

from __future__ import annotations

import html
import re
import struct
import time
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple
from xml.etree import ElementTree

from scripts.document_parsing.core import Block, ParseResult, SourceDocument

from .base import (
    AdapterContext,
    clean_text,
    error_result,
    make_attempt,
    make_block,
    make_result,
    parser_name,
    parser_version,
    source_document_id,
    source_path,
    unavailable_result,
)
from .sniff import sniff_source


def parse_native_hwp(
    source: SourceDocument,
    context: Optional[AdapterContext] = None,
) -> ParseResult:
    parser = parser_name(context, "native-hwp")
    started = time.monotonic()
    path = source_path(source)
    sniffed = sniff_source(path)
    if sniffed.format == "hwpx":
        hwpx_context = AdapterContext(
            parser=parser,
            timeout_seconds=context.timeout_seconds if context else 120,
            env=context.env if context else None,
            raw_output_dir=context.raw_output_dir if context else None,
            options=context.options if context else {},
        )
        return parse_native_hwpx(source, hwpx_context)
    try:
        import olefile
    except ImportError:
        return unavailable_result(source, parser, "optional dependency 'olefile' is not installed")
    try:
        if not olefile.isOleFile(str(path)):
            raise ValueError("input is not an OLE HWP document")
        paragraphs: List[str] = []
        section_count = 0
        with olefile.OleFileIO(str(path)) as document:
            stream_names = ["/".join(parts) for parts in document.listdir(streams=True)]
            compressed = _hwp_is_compressed(document)
            sections = sorted(
                (name for name in stream_names if name.startswith("BodyText/Section")),
                key=_hwp_section_sort_key,
            )
            for section in sections:
                raw = document.openstream(section).read()
                if compressed:
                    import zlib

                    try:
                        raw = zlib.decompress(raw, -15)
                    except zlib.error:
                        raw = zlib.decompress(raw)
                values = _hwp_record_paragraphs(raw)
                if values:
                    section_count += 1
                    paragraphs.extend(values)
            if not paragraphs and "PrvText" in stream_names:
                raw = document.openstream("PrvText").read()
                preview = _decode_preview(raw)
                if preview:
                    paragraphs.extend(_paragraphs(preview))

        blocks = [
            make_block(source, parser, index, "paragraph", paragraph)
            for index, paragraph in enumerate(paragraphs)
            if clean_text(paragraph)
        ]
        return make_result(
            source,
            blocks=blocks,
            attempts=[
                make_attempt(
                    parser,
                    "success",
                    elapsed_ms=int((time.monotonic() - started) * 1000),
                    metadata={"section_count": section_count},
                )
            ],
        )
    except Exception as exc:
        return error_result(source, parser, exc, int((time.monotonic() - started) * 1000))


def parse_native_hwpx(
    source: SourceDocument,
    context: Optional[AdapterContext] = None,
) -> ParseResult:
    parser = parser_name(context, "native-hwpx")
    started = time.monotonic()
    try:
        path = source_path(source)
        if not zipfile.is_zipfile(str(path)):
            raise ValueError("input is not an HWPX ZIP container")
        blocks: List[Block] = []
        table_counter = [0]
        part_count = 0
        with zipfile.ZipFile(str(path)) as archive:
            names = sorted(
                name
                for name in archive.namelist()
                if name.lower().endswith(".xml")
                and (
                    name.lower().startswith("contents/section")
                    or name.lower().startswith("bodytext/section")
                )
            )
            for name in names:
                try:
                    root = ElementTree.fromstring(archive.read(name))
                except (ElementTree.ParseError, KeyError):
                    continue
                part_count += 1
                _walk_hwpx(
                    root,
                    source=source,
                    parser=parser,
                    blocks=blocks,
                    table_counter=table_counter,
                )
        return make_result(
            source,
            blocks=blocks,
            attempts=[
                make_attempt(
                    parser,
                    "success",
                    elapsed_ms=int((time.monotonic() - started) * 1000),
                    metadata={
                        "xml_parts": part_count,
                        "table_count": table_counter[0],
                    },
                )
            ],
        )
    except Exception as exc:
        return error_result(source, parser, exc, int((time.monotonic() - started) * 1000))


def parse_native_html(
    source: SourceDocument,
    context: Optional[AdapterContext] = None,
) -> ParseResult:
    parser = parser_name(context, "native-html")
    started = time.monotonic()
    try:
        path = source_path(source)
        raw = path.read_bytes()
        value = _decode_html(raw)
        document_parser = _SemanticHTMLParser()
        document_parser.feed(value)
        document_parser.close()
        blocks: List[Block] = []
        headings: List[str] = []
        table_counter = 0
        for item_type, payload in document_parser.events:
            if item_type.startswith("h") and len(item_type) == 2 and item_type[1].isdigit():
                level = max(1, min(6, int(item_type[1])))
                text = clean_text(str(payload))
                if not text:
                    continue
                headings = headings[: level - 1]
                headings.append(text)
                blocks.append(
                    make_block(
                        source,
                        parser,
                        len(blocks),
                        "heading",
                        text,
                        section_path=headings,
                    )
                )
            elif item_type in {"p", "li", "caption"}:
                text = clean_text(str(payload))
                if not text:
                    continue
                type_name = {"p": "paragraph", "li": "list_item", "caption": "caption"}[item_type]
                blocks.append(
                    make_block(
                        source,
                        parser,
                        len(blocks),
                        type_name,
                        text,
                        section_path=headings or None,
                    )
                )
            elif item_type == "table":
                rows = _normalize_rows(payload)
                if not rows:
                    continue
                table_id = "{}:{}:t{:04d}".format(
                    source_document_id(source), _safe_id(parser), table_counter
                )
                table_counter += 1
                blocks.append(
                    make_block(
                        source,
                        parser,
                        len(blocks),
                        "table",
                        "\n".join("\t".join(row) for row in rows),
                        section_path=headings or None,
                        table_id=table_id,
                    )
                )
                for row_number, row in enumerate(rows):
                    for column_number, text in enumerate(row):
                        blocks.append(
                            make_block(
                                source,
                                parser,
                                len(blocks),
                                "table_cell",
                                text,
                                section_path=headings or None,
                                table_id=table_id,
                                row=row_number,
                                column=column_number,
                            )
                        )

        used_trafilatura = False
        selected_parser = parser
        if sum(len(getattr(block, "text", "")) for block in blocks) < 20:
            allow_fallback = not context or bool(context.options.get("trafilatura_fallback", True))
            if allow_fallback:
                try:
                    import trafilatura

                    extracted = trafilatura.extract(value) or ""
                    if extracted.strip():
                        used_trafilatura = True
                        selected_parser = (
                            "{}/trafilatura".format(parser.rsplit("/", 1)[0])
                            if "/" in parser
                            else "{}-trafilatura".format(parser)
                        )
                        blocks = [
                            make_block(
                                source,
                                selected_parser,
                                index,
                                "paragraph",
                                paragraph,
                            )
                            for index, paragraph in enumerate(_paragraphs(extracted))
                        ]
                except ImportError:
                    pass
        return make_result(
            source,
            blocks=blocks,
            attempts=[
                make_attempt(
                    selected_parser,
                    "success",
                    elapsed_ms=int((time.monotonic() - started) * 1000),
                    metadata={
                        "table_count": table_counter,
                        "trafilatura_fallback": used_trafilatura,
                        "fallback_from": parser
                        if used_trafilatura
                        else None,
                    },
                    parser_version=(
                        str(context.options.get("trafilatura_version"))
                        if (
                            used_trafilatura
                            and context
                            and context.options.get("trafilatura_version")
                        )
                        else None
                    ),
                )
            ],
        )
    except Exception as exc:
        return error_result(source, parser, exc, int((time.monotonic() - started) * 1000))


def parse_native_office(
    source: SourceDocument,
    context: Optional[AdapterContext] = None,
) -> ParseResult:
    format_name = sniff_source(source_path(source)).format
    if format_name == "docx":
        return _parse_docx(source, context)
    if format_name == "xlsx":
        return _parse_xlsx(source, context)
    if format_name == "xls":
        return _parse_xls(source, context)
    if format_name == "pptx":
        return _parse_pptx(source, context)
    parser = parser_name(context, "native-office")
    return make_result(
        source,
        attempts=[
            make_attempt(
                parser,
                "unsupported",
                reason="unsupported Office format: {}".format(format_name),
            )
        ],
    )


def _parse_docx(
    source: SourceDocument,
    context: Optional[AdapterContext],
) -> ParseResult:
    parser = parser_name(context, "python-docx")
    started = time.monotonic()
    try:
        import docx
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        return unavailable_result(source, parser, "optional dependency 'python-docx' is not installed")
    try:
        document = docx.Document(str(source_path(source)))
        blocks: List[Block] = []
        headings: List[str] = []
        table_number = 0
        for child in document.element.body.iterchildren():
            if isinstance(child, CT_P):
                paragraph = Paragraph(child, document)
                text = clean_text(paragraph.text)
                if not text:
                    continue
                style_name = str(
                    getattr(getattr(paragraph, "style", None), "name", "")
                    or ""
                )
                match = re.match(
                    r"Heading\s+([1-6])",
                    style_name,
                    flags=re.IGNORECASE,
                )
                if match:
                    level = int(match.group(1))
                    headings = headings[: level - 1]
                    headings.append(text)
                    block_type = "heading"
                else:
                    block_type = (
                        "list_item"
                        if "list" in style_name.lower()
                        else "paragraph"
                    )
                blocks.append(
                    make_block(
                        source,
                        parser,
                        len(blocks),
                        block_type,
                        text,
                        section_path=headings or None,
                    )
                )
            elif isinstance(child, CT_Tbl):
                table = Table(child, document)
                rows = _normalize_rows(
                    [
                        [cell.text for cell in row.cells]
                        for row in table.rows
                    ]
                )
                if rows:
                    _append_table_blocks(
                        source,
                        parser,
                        blocks,
                        rows,
                        table_number,
                        headings,
                    )
                table_number += 1
        return make_result(
            source,
            blocks=blocks,
            attempts=[
                make_attempt(
                    parser,
                    "success",
                    elapsed_ms=int((time.monotonic() - started) * 1000),
                    metadata={"table_count": len(document.tables)},
                    parser_version=parser_version(context),
                )
            ],
        )
    except Exception as exc:
        return error_result(source, parser, exc, int((time.monotonic() - started) * 1000))


def _parse_xlsx(
    source: SourceDocument,
    context: Optional[AdapterContext],
) -> ParseResult:
    parser = parser_name(context, "openpyxl")
    started = time.monotonic()
    try:
        import openpyxl
    except ImportError:
        return unavailable_result(source, parser, "optional dependency 'openpyxl' is not installed")
    try:
        workbook = openpyxl.load_workbook(
            str(source_path(source)), read_only=True, data_only=True
        )
        blocks: List[Block] = []
        try:
            for table_number, sheet in enumerate(workbook.worksheets):
                heading = clean_text(sheet.title)
                blocks.append(
                    make_block(
                        source,
                        parser,
                        len(blocks),
                        "heading",
                        heading,
                        section_path=[heading],
                    )
                )
                rows = _normalize_rows(
                    [
                        ["" if value is None else str(value) for value in row]
                        for row in sheet.iter_rows(values_only=True)
                    ]
                )
                if rows:
                    _append_table_blocks(source, parser, blocks, rows, table_number, [heading])
        finally:
            workbook.close()
        return make_result(
            source,
            blocks=blocks,
            attempts=[
                make_attempt(
                    parser,
                    "success",
                    elapsed_ms=int((time.monotonic() - started) * 1000),
                    metadata={"sheet_count": len(workbook.sheetnames)},
                    parser_version=parser_version(context),
                )
            ],
        )
    except Exception as exc:
        return error_result(source, parser, exc, int((time.monotonic() - started) * 1000))


def _parse_xls(
    source: SourceDocument,
    context: Optional[AdapterContext],
) -> ParseResult:
    parser = parser_name(context, "xlrd")
    started = time.monotonic()
    try:
        import xlrd
    except ImportError:
        return unavailable_result(source, parser, "optional dependency 'xlrd' is not installed")
    try:
        workbook = xlrd.open_workbook(str(source_path(source)), on_demand=True)
        blocks: List[Block] = []
        for table_number, sheet_name in enumerate(workbook.sheet_names()):
            sheet = workbook.sheet_by_name(sheet_name)
            heading = clean_text(sheet_name)
            blocks.append(
                make_block(
                    source,
                    parser,
                    len(blocks),
                    "heading",
                    heading,
                    section_path=[heading],
                )
            )
            rows = _normalize_rows(
                [[str(sheet.cell_value(row, column)) for column in range(sheet.ncols)]
                 for row in range(sheet.nrows)]
            )
            if rows:
                _append_table_blocks(source, parser, blocks, rows, table_number, [heading])
        workbook.release_resources()
        return make_result(
            source,
            blocks=blocks,
            attempts=[
                make_attempt(
                    parser,
                    "success",
                    elapsed_ms=int((time.monotonic() - started) * 1000),
                    metadata={"sheet_count": len(workbook.sheet_names())},
                    parser_version=parser_version(context),
                )
            ],
        )
    except Exception as exc:
        return error_result(source, parser, exc, int((time.monotonic() - started) * 1000))


def _parse_pptx(
    source: SourceDocument,
    context: Optional[AdapterContext],
) -> ParseResult:
    parser = parser_name(context, "pptx-xml")
    started = time.monotonic()
    try:
        blocks: List[Block] = []
        with zipfile.ZipFile(str(source_path(source))) as archive:
            names = sorted(
                (
                    name
                    for name in archive.namelist()
                    if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
                ),
                key=_number_in_name,
            )
            for page_number, name in enumerate(names, start=1):
                root = ElementTree.fromstring(archive.read(name))
                texts = [
                    clean_text(element.text or "")
                    for element in root.iter()
                    if _local_name(element.tag) == "t" and clean_text(element.text or "")
                ]
                for paragraph in texts:
                    blocks.append(
                        make_block(
                            source,
                            parser,
                            len(blocks),
                            "paragraph",
                            paragraph,
                            page=page_number,
                        )
                    )
        return make_result(
            source,
            blocks=blocks,
            attempts=[
                make_attempt(
                    parser,
                    "success",
                    elapsed_ms=int((time.monotonic() - started) * 1000),
                    metadata={"slide_count": len(names)},
                    parser_version=parser_version(context),
                )
            ],
        )
    except Exception as exc:
        return error_result(source, parser, exc, int((time.monotonic() - started) * 1000))


def _walk_hwpx(
    element: ElementTree.Element,
    source: SourceDocument,
    parser: str,
    blocks: List[Block],
    table_counter: List[int],
) -> None:
    for child in list(element):
        local = _local_name(child.tag)
        if local in {"tbl", "table"}:
            rows = _hwpx_table_rows(child)
            if rows:
                number = table_counter[0]
                table_counter[0] += 1
                _append_table_blocks(source, parser, blocks, rows, number, None)
            continue
        if local in {"p", "paragraph"}:
            text = _element_text_without_tables(child)
            if text:
                blocks.append(
                    make_block(source, parser, len(blocks), "paragraph", text)
                )
            for table in _top_level_descendant_tables(child):
                rows = _hwpx_table_rows(table)
                if rows:
                    number = table_counter[0]
                    table_counter[0] += 1
                    _append_table_blocks(
                        source, parser, blocks, rows, number, None
                    )
            continue
        _walk_hwpx(child, source, parser, blocks, table_counter)


def _hwpx_table_rows(table: ElementTree.Element) -> List[List[str]]:
    cells: Dict[Tuple[int, int], str] = {}
    max_row = -1
    max_column = -1
    for row_index, row in enumerate(
        element for element in table.iter() if _local_name(element.tag) in {"tr", "row"}
    ):
        sequential_column = 0
        for cell in (
            element for element in list(row) if _local_name(element.tag) in {"tc", "cell"}
        ):
            attrs = {_local_name(key): value for key, value in cell.attrib.items()}
            cell_row = int(attrs.get("rowAddr", attrs.get("row", row_index)) or row_index)
            cell_column = int(
                attrs.get("colAddr", attrs.get("column", sequential_column)) or sequential_column
            )
            row_span = max(1, int(attrs.get("rowSpan", 1) or 1))
            column_span = max(1, int(attrs.get("colSpan", 1) or 1))
            cells[(cell_row, cell_column)] = _element_text(cell)
            for span_row in range(cell_row, cell_row + row_span):
                for span_column in range(cell_column, cell_column + column_span):
                    cells.setdefault((span_row, span_column), "")
            max_row = max(max_row, cell_row + row_span - 1)
            max_column = max(max_column, cell_column + column_span - 1)
            sequential_column = cell_column + column_span
    if max_row < 0 or max_column < 0:
        return []
    return [
        [cells.get((row, column), "") for column in range(max_column + 1)]
        for row in range(max_row + 1)
    ]


def _append_table_blocks(
    source: SourceDocument,
    parser: str,
    blocks: List[Block],
    rows: Sequence[Sequence[str]],
    table_number: int,
    section_path: Optional[Sequence[str]],
) -> None:
    normalized = _normalize_rows(rows)
    table_id = "{}:{}:t{:04d}".format(
        source_document_id(source), _safe_id(parser), table_number
    )
    blocks.append(
        make_block(
            source,
            parser,
            len(blocks),
            "table",
            "\n".join("\t".join(row) for row in normalized),
            section_path=section_path,
            table_id=table_id,
        )
    )
    for row_number, row in enumerate(normalized):
        for column_number, value in enumerate(row):
            blocks.append(
                make_block(
                    source,
                    parser,
                    len(blocks),
                    "table_cell",
                    value,
                    section_path=section_path,
                    table_id=table_id,
                    row=row_number,
                    column=column_number,
                )
            )


class _SemanticHTMLParser(HTMLParser):
    def __init__(self) -> None:
        HTMLParser.__init__(self, convert_charrefs=True)
        self.events: List[Tuple[str, Any]] = []
        self._capture: List[Tuple[str, List[str]]] = []
        self._ignored_depth = 0
        self._table_depth = 0
        self._table_rows: List[List[str]] = []
        self._current_row: Optional[List[str]] = None
        self._current_cell: Optional[List[str]] = None

    def handle_starttag(self, tag: str, attrs: List[Tuple[str, Optional[str]]]) -> None:
        tag = tag.lower()
        if tag in {"script", "style", "noscript", "svg"}:
            self._ignored_depth += 1
            return
        if self._ignored_depth:
            return
        if tag == "table":
            self._table_depth += 1
            if self._table_depth == 1:
                self._table_rows = []
            return
        if self._table_depth:
            if tag == "tr":
                self._current_row = []
            elif tag in {"td", "th"}:
                self._current_cell = []
            return
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "figcaption", "caption", "title"}:
            mapped = "h1" if tag == "title" else "caption" if tag == "figcaption" else tag
            self._capture.append((mapped, []))

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in {"script", "style", "noscript", "svg"}:
            if self._ignored_depth:
                self._ignored_depth -= 1
            return
        if self._ignored_depth:
            return
        if self._table_depth:
            if tag in {"td", "th"} and self._current_cell is not None:
                if self._current_row is None:
                    self._current_row = []
                self._current_row.append(clean_text("".join(self._current_cell)))
                self._current_cell = None
            elif tag == "tr" and self._current_row is not None:
                self._table_rows.append(self._current_row)
                self._current_row = None
            elif tag == "table":
                self._table_depth -= 1
                if self._table_depth == 0:
                    self.events.append(("table", list(self._table_rows)))
                    self._table_rows = []
            return
        mapped = "h1" if tag == "title" else "caption" if tag == "figcaption" else tag
        for index in range(len(self._capture) - 1, -1, -1):
            captured_tag, parts = self._capture[index]
            if captured_tag == mapped:
                del self._capture[index]
                text = clean_text("".join(parts))
                if text:
                    self.events.append((captured_tag, text))
                break

    def handle_data(self, data: str) -> None:
        if self._ignored_depth:
            return
        if self._current_cell is not None:
            self._current_cell.append(data)
            return
        if self._capture:
            self._capture[-1][1].append(data)


def _hwp_is_compressed(document: Any) -> bool:
    header = document.openstream("FileHeader").read()
    return len(header) >= 40 and bool(struct.unpack_from("<I", header, 36)[0] & 0x01)


def _hwp_section_sort_key(name: str) -> Tuple[int, str]:
    match = re.search(r"Section(\d+)$", name)
    return (int(match.group(1)), name) if match else (10 ** 9, name)


def _hwp_record_paragraphs(data: bytes) -> List[str]:
    result: List[str] = []
    offset = 0
    while offset + 4 <= len(data):
        header = int.from_bytes(data[offset : offset + 4], "little")
        offset += 4
        tag_id = header & 0x3FF
        size = (header >> 20) & 0xFFF
        if size == 0xFFF:
            if offset + 4 > len(data):
                break
            size = int.from_bytes(data[offset : offset + 4], "little")
            offset += 4
        if offset + size > len(data):
            break
        payload = data[offset : offset + size]
        offset += size
        if tag_id == 67:
            text = clean_text(payload.decode("utf-16le", errors="ignore").replace("\r", "\n"))
            if text:
                result.extend(_paragraphs(text))
    return result


def _decode_preview(data: bytes) -> str:
    candidates = [
        clean_text(data.decode(encoding, errors="ignore"))
        for encoding in ("utf-16le", "cp949", "utf-8")
    ]
    return max(candidates, key=len, default="")


def _element_text(element: ElementTree.Element) -> str:
    parts: List[str] = []
    for item in element.iter():
        local = _local_name(item.tag)
        if local in {"t", "text"} and item.text:
            parts.append(item.text)
    if not parts:
        parts = [text for text in element.itertext() if text]
    return clean_text(" ".join(parts))


def _element_text_without_tables(element: ElementTree.Element) -> str:
    parts: List[str] = []

    def walk(item: ElementTree.Element) -> None:
        if _local_name(item.tag) in {"tbl", "table"}:
            return
        if _local_name(item.tag) in {"t", "text"} and item.text:
            parts.append(item.text)
        for child in list(item):
            walk(child)

    walk(element)
    return clean_text(" ".join(parts))


def _top_level_descendant_tables(
    element: ElementTree.Element,
) -> Iterable[ElementTree.Element]:
    for child in list(element):
        if _local_name(child.tag) in {"tbl", "table"}:
            yield child
        else:
            for nested in _top_level_descendant_tables(child):
                yield nested


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1].rsplit(":", 1)[-1]


def _normalize_rows(rows: Sequence[Sequence[Any]]) -> List[List[str]]:
    width = max((len(row) for row in rows), default=0)
    result: List[List[str]] = []
    for raw in rows:
        row = [clean_text("" if value is None else str(value)) for value in raw]
        row.extend([""] * (width - len(row)))
        if any(row):
            result.append(row)
    return result


def _paragraphs(value: str) -> List[str]:
    cleaned = clean_text(value)
    if not cleaned:
        return []
    chunks = re.split(r"\n\s*\n+", cleaned)
    if len(chunks) == 1:
        chunks = cleaned.splitlines()
    return [clean_text(chunk) for chunk in chunks if clean_text(chunk)]


def _decode_html(raw: bytes) -> str:
    prefix = raw[:4096].decode("ascii", errors="ignore")
    match = re.search(r"charset\s*=\s*[\"']?([A-Za-z0-9._-]+)", prefix, re.IGNORECASE)
    encodings = [match.group(1)] if match else []
    encodings.extend(["utf-8", "cp949"])
    for encoding in encodings:
        try:
            return html.unescape(raw.decode(encoding))
        except (LookupError, UnicodeDecodeError):
            continue
    return html.unescape(raw.decode("utf-8", errors="replace"))


def _safe_id(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z._-]+", "-", value).strip("-") or "parser"


def _number_in_name(name: str) -> int:
    match = re.search(r"(\d+)", Path(name).stem)
    return int(match.group(1)) if match else 10 ** 9
