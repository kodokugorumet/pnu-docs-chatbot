from __future__ import annotations

import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from scripts.document_parsing.adapters import (
    AdapterContext,
    classify_pdf_page,
    overlay_pdf_tables,
    parse_native_html,
    parse_native_hwpx,
    parse_native_office,
    parse_subprocess,
    parse_tesseract,
    parse_unhwp,
    sniff_source,
)
from scripts.document_parsing.adapters.sniff import OLE_MAGIC
from scripts.document_parsing.core import Block, ParseResult, SourceDocument


def make_source(path: Path, document_id: str = "doc_test") -> SourceDocument:
    return SourceDocument(
        path=path,
        document_id=document_id,
        relative_path=path.name,
        size_bytes=path.stat().st_size,
    )


class SniffSourceTests(unittest.TestCase):
    def test_pdf_magic_wins_over_extension(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "wrong.hwp"
            path.write_bytes(b"%PDF-1.7\nfixture")

            result = sniff_source(path)

            self.assertEqual(result.format, "pdf")
            self.assertEqual(result.reason, "pdf_magic")

    def test_hwpx_zip_is_detected_when_renamed_hwp(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "renamed.hwp"
            with zipfile.ZipFile(str(path), "w") as archive:
                archive.writestr("mimetype", "application/hwp+zip")
                archive.writestr("version.xml", "<version/>")
                archive.writestr("Contents/section0.xml", "<section/>")

            result = sniff_source(path)

            self.assertEqual(result.format, "hwpx")
            self.assertEqual(result.container, "zip")

    def test_fasoo_drm_is_not_misrouted_by_hwp_extension(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "protected.hwp"
            path.write_bytes(b"\x9b DRMONE  This Document is encrypted by Fasoo DRM")

            result = sniff_source(path)

            self.assertEqual(result.format, "drm")
            self.assertEqual(result.reason, "fasoo_drm_signature")

    def test_ole_xls_is_routed_to_the_legacy_excel_adapter(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "legacy.xls"
            path.write_bytes(OLE_MAGIC + b"\x00" * 64)

            result = sniff_source(path)

            self.assertEqual(result.format, "xls")
            self.assertEqual(result.mime_type, "application/vnd.ms-excel")
            self.assertEqual(result.reason, "ole_magic")


class PageClassificationTests(unittest.TestCase):
    def test_scan_has_precedence(self) -> None:
        result = classify_pdf_page(
            {
                "text_char_count": 15,
                "word_count": 2,
                "image_area_ratio": 0.8,
                "has_tables": True,
                "column_count": 2,
            }
        )
        self.assertEqual(result, "scan")

    def test_table_then_complex_then_clean(self) -> None:
        base = {
            "text_char_count": 200,
            "word_count": 40,
            "image_area_ratio": 0.0,
            "reading_order_anomaly": False,
        }
        self.assertEqual(classify_pdf_page(dict(base, has_tables=True, column_count=2)), "table")
        self.assertEqual(classify_pdf_page(dict(base, has_tables=False, column_count=2)), "complex")
        self.assertEqual(classify_pdf_page(dict(base, has_tables=False, column_count=1)), "clean")


class NativeAdapterTests(unittest.TestCase):
    def test_hwpx_emits_paragraph_and_table_parent_cells(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.hwpx"
            section = """\
<hs:section xmlns:hs="urn:section" xmlns:hp="urn:paragraph">
  <hp:p><hp:run><hp:t>첫 문단</hp:t></hp:run></hp:p>
  <hp:tbl>
    <hp:tr>
      <hp:tc rowAddr="0" colAddr="0"><hp:p><hp:t>항목</hp:t></hp:p></hp:tc>
      <hp:tc rowAddr="0" colAddr="1"><hp:p><hp:t>값</hp:t></hp:p></hp:tc>
    </hp:tr>
  </hp:tbl>
</hs:section>
"""
            with zipfile.ZipFile(str(path), "w") as archive:
                archive.writestr("mimetype", "application/hwp+zip")
                archive.writestr("version.xml", "<version/>")
                archive.writestr("Contents/section0.xml", section)

            result = parse_native_hwpx(make_source(path))

            self.assertEqual(result.attempts[0].status, "success")
            self.assertEqual(
                [block.block_type for block in result.blocks],
                ["paragraph", "table", "table_cell", "table_cell"],
            )
            self.assertEqual(result.blocks[1].text, "항목\t값")
            self.assertEqual(result.blocks[2].row, 0)
            self.assertEqual(result.blocks[3].column, 1)
            result.validate()

    def test_hwpx_finds_table_nested_in_paragraph_control(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "nested.hwpx"
            section = """\
<hs:section xmlns:hs="urn:section" xmlns:hp="urn:paragraph">
  <hp:p><hp:run><hp:t>표 앞 문장</hp:t><hp:ctrl><hp:tbl><hp:tr>
    <hp:tc><hp:p><hp:t>A</hp:t></hp:p></hp:tc>
    <hp:tc><hp:p><hp:t>B</hp:t></hp:p></hp:tc>
  </hp:tr></hp:tbl></hp:ctrl></hp:run></hp:p>
</hs:section>
"""
            with zipfile.ZipFile(str(path), "w") as archive:
                archive.writestr("mimetype", "application/hwp+zip")
                archive.writestr("version.xml", "<version/>")
                archive.writestr("Contents/section0.xml", section)

            result = parse_native_hwpx(make_source(path))

            self.assertEqual(
                [block.block_type for block in result.blocks],
                ["paragraph", "table", "table_cell", "table_cell"],
            )
            self.assertEqual(result.blocks[0].text, "표 앞 문장")
            self.assertEqual(result.blocks[1].text, "A\tB")

    def test_html_preserves_headings_lists_and_tables(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.html"
            path.write_text(
                """
<!doctype html><html><body>
<h1>안내</h1><p>본문 내용입니다.</p><ul><li>첫 항목</li></ul>
<table><tr><th>이름</th><th>값</th></tr><tr><td>A</td><td>1</td></tr></table>
</body></html>
""",
                encoding="utf-8",
            )

            result = parse_native_html(make_source(path))

            types = [block.block_type for block in result.blocks]
            self.assertEqual(types[:3], ["heading", "paragraph", "list_item"])
            self.assertEqual(types.count("table"), 1)
            self.assertEqual(types.count("table_cell"), 4)
            self.assertEqual(result.blocks[1].section_path, ("안내",))
            result.validate()

    def test_html_trafilatura_fallback_records_its_own_provenance(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "empty.html"
            path.write_text("<html><body></body></html>", encoding="utf-8")
            trafilatura = mock.Mock()
            trafilatura.extract.return_value = (
                "fallback로 추출한 충분히 긴 본문입니다."
            )
            with mock.patch.dict(
                sys.modules,
                {"trafilatura": trafilatura},
            ):
                result = parse_native_html(
                    make_source(path),
                    AdapterContext(
                        parser="baseline/html-dom",
                        options={
                            "trafilatura_fallback": True,
                            "trafilatura_version": "2.0.0",
                        },
                    ),
                )

            self.assertEqual(
                {block.parser for block in result.blocks},
                {"baseline/trafilatura"},
            )
            self.assertEqual(
                result.attempts[0].parser,
                "baseline/trafilatura",
            )
            self.assertEqual(
                result.attempts[0].parser_version,
                "2.0.0",
            )

    def test_docx_preserves_paragraph_table_paragraph_body_order(self) -> None:
        try:
            import docx
        except ImportError:
            self.skipTest("python-docx is not installed")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ordered.docx"
            document = docx.Document()
            document.add_paragraph("표 앞 문단")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "항목"
            table.cell(0, 1).text = "값"
            document.add_paragraph("표 뒤 문단")
            document.save(str(path))

            result = parse_native_office(make_source(path))

            self.assertEqual(result.attempts[0].status, "success")
            self.assertEqual(
                [block.block_type for block in result.blocks],
                [
                    "paragraph",
                    "table",
                    "table_cell",
                    "table_cell",
                    "paragraph",
                ],
            )
            self.assertEqual(result.blocks[0].text, "표 앞 문단")
            self.assertEqual(result.blocks[-1].text, "표 뒤 문단")

    def test_xlsx_stops_while_streaming_when_block_limit_is_exceeded(self) -> None:
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is not installed")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "large.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["항목", "값"])
            sheet.append(["A", "1"])
            workbook.save(str(path))
            workbook.close()

            result = parse_native_office(
                make_source(path),
                AdapterContext(
                    parser="baseline/openpyxl",
                    options={"max_blocks": 4},
                ),
            )

            self.assertEqual(result.blocks, [])
            self.assertEqual(result.attempts[0].status, "error")
            self.assertIn(
                "block_limit_exceeded",
                result.attempts[0].reason,
            )


class SubprocessAdapterTests(unittest.TestCase):
    def test_missing_command_is_explicitly_unavailable(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.txt"
            path.write_text("input", encoding="utf-8")

            result = parse_subprocess(
                make_source(path),
                AdapterContext(parser="worker"),
                command_env="PARSER_TEST_COMMAND_THAT_DOES_NOT_EXIST",
            )

            self.assertEqual(result.attempts[0].status, "unavailable")
            self.assertEqual(result.blocks, [])

    def test_shell_worker_commands_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.txt"
            path.write_text("input", encoding="utf-8")

            result = parse_subprocess(
                make_source(path),
                AdapterContext(
                    parser="unsafe-worker",
                    command=("sh", "-c", "tool {input}"),
                ),
            )

            self.assertEqual(result.attempts[0].status, "unavailable")
            self.assertIn(
                "shell launchers",
                result.attempts[0].reason,
            )

    def test_pdf_tesseract_reports_missing_renderer(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "scan.pdf"
            path.write_bytes(b"%PDF-1.7\nfixture")

            original_import = __import__

            def missing_fitz(name, *args, **kwargs):
                if name == "fitz":
                    raise ImportError("fixture")
                return original_import(name, *args, **kwargs)

            with mock.patch("builtins.__import__", side_effect=missing_fitz):
                result = parse_tesseract(make_source(path))

            self.assertEqual(result.attempts[0].status, "unavailable")
            self.assertIn("PyMuPDF", result.attempts[0].reason)

    def test_json_worker_output_is_canonicalized(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.txt"
            path.write_text("input", encoding="utf-8")
            payload = {
                "blocks": [
                    {"block_type": "heading", "text": "제목", "page": 1},
                    {"block_type": "paragraph", "text": "본문", "page": 1},
                ],
                "metadata": {"worker": "fixture"},
            }
            command = (
                sys.executable,
                "-c",
                "import json; print(json.dumps({!r}, ensure_ascii=False))".format(payload),
            )

            result = parse_subprocess(
                make_source(path),
                AdapterContext(parser="fixture-worker", command=command),
            )

            self.assertEqual(result.attempts[0].status, "success")
            self.assertEqual(result.attempts[0].metadata["worker"], "fixture")
            self.assertEqual([block.reading_order for block in result.blocks], [0, 1])
            self.assertEqual(result.blocks[0].page, 1)
            result.validate()

    def test_tagged_text_worker_output_sets_page(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.txt"
            path.write_text("input", encoding="utf-8")
            script = "print('[page 2]\\n[heading] 제목\\n[paragraph] 본문')"

            result = parse_subprocess(
                make_source(path),
                AdapterContext(
                    parser="tagged-worker",
                    command=(sys.executable, "-c", script),
                ),
            )

            self.assertEqual([block.block_type for block in result.blocks], ["heading", "paragraph"])
            self.assertTrue(all(block.page == 2 for block in result.blocks))

    def test_anonymous_worker_table_cells_share_parent_table_id(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.txt"
            path.write_text("input", encoding="utf-8")
            payload = {
                "blocks": [
                    {"block_type": "table", "text": "A\tB"},
                    {"block_type": "table_cell", "text": "A", "row": 0, "column": 0},
                    {"block_type": "table_cell", "text": "B", "row": 0, "column": 1},
                ]
            }
            command = (
                sys.executable,
                "-c",
                "import json; print(json.dumps({!r}))".format(payload),
            )

            result = parse_subprocess(
                make_source(path),
                AdapterContext(parser="table-worker", command=command),
            )

            self.assertEqual(len({block.table_id for block in result.blocks}), 1)
            self.assertEqual(result.blocks[0].text, "A\tB")
            result.validate()

    def test_empty_worker_table_parent_is_preserved(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.txt"
            path.write_text("input", encoding="utf-8")
            payload = {
                "blocks": [
                    {
                        "block_type": "table",
                        "text": "",
                        "table_id": "layout-grid",
                    },
                    {
                        "block_type": "table_cell",
                        "text": "",
                        "table_id": "layout-grid",
                        "row": 0,
                        "column": 0,
                    },
                ]
            }
            command = (
                sys.executable,
                "-c",
                "import json; print(json.dumps({!r}))".format(payload),
            )

            result = parse_subprocess(
                make_source(path),
                AdapterContext(parser="table-worker", command=command),
            )

            self.assertEqual(
                [block.block_type for block in result.blocks],
                ["table", "table_cell"],
            )
            self.assertEqual(result.blocks[0].text, "")
            self.assertEqual(
                result.blocks[0].table_id,
                result.blocks[1].table_id,
            )
            result.validate()

    def test_per_page_worker_raw_outputs_do_not_overwrite_each_other(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            input_path = root / "input.txt"
            input_path.write_text("input", encoding="utf-8")
            raw_root = root / "raw"
            command = (
                sys.executable,
                "-c",
                "print('[paragraph] page output')",
            )
            source = make_source(input_path)

            for page in (1, 2):
                page_source = SourceDocument(
                    path=source.path,
                    document_id=source.document_id,
                    relative_path="{}#page={}".format(
                        source.relative_path, page
                    ),
                    size_bytes=source.size_bytes,
                )
                result = parse_subprocess(
                    page_source,
                    AdapterContext(
                        parser="page-worker",
                        command=command,
                        raw_output_dir=raw_root,
                    ),
                )
                self.assertEqual(result.attempts[0].status, "success")

            output_directories = [
                path for path in raw_root.iterdir() if path.is_dir()
            ]
            self.assertEqual(len(output_directories), 2)
            self.assertTrue(
                all(
                    (path / "worker.stdout").is_file()
                    for path in output_directories
                )
            )

    def test_unhwp_raw_content_sections_are_translated(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.hwp"
            path.write_text("fixture", encoding="utf-8")
            payload = {
                "metadata": {"title": "fixture"},
                "sections": [
                    {
                        "index": 0,
                        "content": [
                            {
                                "Paragraph": {
                                    "style": {"heading_level": 1},
                                    "content": [{"Text": {"text": "제목", "style": {}}}],
                                }
                            },
                            {
                                "Paragraph": {
                                    "style": {"heading_level": 0},
                                    "content": [
                                        {
                                            "Text": {
                                                "text": "일반 본문",
                                                "style": {},
                                            }
                                        }
                                    ],
                                }
                            },
                            {
                                "Paragraph": {
                                    "style": {
                                        "heading_level": 4,
                                        "list_style": {"kind": "bullet"},
                                    },
                                    "content": [
                                        {
                                            "Text": {
                                                "text": "◦ 예산조치 사항",
                                                "style": {},
                                            }
                                        }
                                    ],
                                }
                            },
                            {
                                "Table": {
                                    "rows": [
                                        {
                                            "cells": [
                                                {
                                                    "content": [
                                                        {
                                                            "style": {},
                                                            "content": [
                                                                {"Text": {"text": "셀", "style": {}}}
                                                            ],
                                                        }
                                                    ],
                                                    "rowspan": 1,
                                                    "colspan": 1,
                                                }
                                            ]
                                        }
                                    ]
                                }
                            },
                        ],
                    }
                ],
            }
            command = (
                sys.executable,
                "-c",
                "import json; print(json.dumps({!r}, ensure_ascii=False))".format(payload),
            )

            result = parse_subprocess(
                make_source(path),
                AdapterContext(parser="unhwp", command=command),
            )

            self.assertEqual(
                [block.block_type for block in result.blocks],
                [
                    "heading",
                    "paragraph",
                    "list_item",
                    "table",
                    "table_cell",
                ],
            )
            self.assertEqual(result.blocks[1].text, "일반 본문")
            self.assertEqual(result.blocks[2].text, "◦ 예산조치 사항")
            self.assertEqual(result.blocks[3].text, "셀")
            self.assertEqual(
                result.attempts[0].metadata["output_format"],
                "unhwp_raw_content",
            )
            result.validate()

    def test_managed_unhwp_does_not_fall_back_to_environment_command(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.hwp"
            path.write_text("fixture", encoding="utf-8")
            with mock.patch.dict(
                "os.environ",
                {"PARSER_UNHWP_CMD": sys.executable},
            ):
                result = parse_unhwp(
                    make_source(path),
                    AdapterContext(
                        parser="managed/unhwp",
                        options={"managed_runtime": True},
                    ),
                )

            self.assertEqual(result.attempts[0].status, "unavailable")
            self.assertIn("no command configured", result.attempts[0].reason)


class OverlayTests(unittest.TestCase):
    def test_overlay_removes_exact_cell_duplicate_and_reindexes(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.pdf"
            path.write_bytes(b"%PDF-fixture")
            source = make_source(path)
            base = ParseResult(
                source,
                blocks=[
                    Block(
                        source.document_id,
                        "old-0",
                        "paragraph",
                        "셀 값",
                        0,
                        "pymupdf",
                        page=1,
                    ),
                    Block(
                        source.document_id,
                        "old-1",
                        "paragraph",
                        "남길 본문",
                        1,
                        "pymupdf",
                        page=1,
                    ),
                ],
            )
            table_id = "doc_test:pdfplumber:t0000"
            tables = ParseResult(
                source,
                blocks=[
                    Block(
                        source.document_id,
                        "table-0",
                        "table",
                        "셀 값\t다른 값",
                        0,
                        "pdfplumber",
                        page=1,
                        table_id=table_id,
                    ),
                    Block(
                        source.document_id,
                        "table-1",
                        "table_cell",
                        "셀 값",
                        1,
                        "pdfplumber",
                        page=1,
                        table_id=table_id,
                        row=0,
                        column=0,
                    ),
                ],
            )

            result = overlay_pdf_tables(source, base, tables, "baseline/pdf")

            self.assertNotIn("셀 값", [block.text for block in result.blocks if block.block_type == "paragraph"])
            self.assertEqual([block.reading_order for block in result.blocks], [0, 1, 2])
            result.validate()

    def test_overlay_keeps_existing_structured_table_as_primary(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.pdf"
            path.write_bytes(b"%PDF-fixture")
            source = make_source(path)
            docling_table_id = "doc_test:docling:t0000"
            base = ParseResult(
                source,
                blocks=[
                    Block(
                        source.document_id,
                        "docling-table",
                        "table",
                        "구조화 표",
                        0,
                        "cascade/docling",
                        page=1,
                        table_id=docling_table_id,
                    ),
                    Block(
                        source.document_id,
                        "docling-cell",
                        "table_cell",
                        "구조화 표",
                        1,
                        "cascade/docling",
                        page=1,
                        table_id=docling_table_id,
                        row=0,
                        column=0,
                    ),
                ],
            )
            fallback_table_id = "doc_test:pdfplumber:t0000"
            tables = ParseResult(
                source,
                blocks=[
                    Block(
                        source.document_id,
                        "fallback-table",
                        "table",
                        "중복 표",
                        0,
                        "cascade/pdfplumber",
                        page=1,
                        table_id=fallback_table_id,
                    ),
                    Block(
                        source.document_id,
                        "fallback-cell",
                        "table_cell",
                        "중복 표",
                        1,
                        "cascade/pdfplumber",
                        page=1,
                        table_id=fallback_table_id,
                        row=0,
                        column=0,
                    ),
                ],
            )

            result = overlay_pdf_tables(
                source, base, tables, "cascade/pdf-composite"
            )

            parents = [
                block
                for block in result.blocks
                if block.block_type == "table"
            ]
            self.assertEqual(len(parents), 1)
            self.assertEqual(parents[0].text, "구조화 표")
            self.assertEqual(parents[0].parser, "cascade/docling")
            result.validate()


if __name__ == "__main__":
    unittest.main()
